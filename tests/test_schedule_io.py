"""
Unit tests for backend.app.services.schedule_io.
Verifies PDF generation, Excel generation, Excel parsing, and safe Content-Disposition headers.
"""

import unittest
from datetime import time
import io
import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from fastapi import HTTPException

from backend.app import models
from backend.app.services.schedule_io import (
    _attachment_header,
    attachment_header,
    build_schedule_pdf,
    build_schedule_excel,
    import_schedules_from_excel,
)


class ScheduleIOTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        models.Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine)
        self.db = self.Session()

        # Seed fixtures
        self.dept = models.Department(id=1, name="College of Arts, Sciences & Technology", code="CAST")
        self.sem = models.Semester(id=1, academic_year="2026-2027", term="1st", is_active=True)
        self.curr = models.Curriculum(id=1, code="CS101", name="Intro to <CS> & Programming", units=3, department_id=1, type="lecture")
        self.fac = models.Faculty(id=1, first_name="Grace", last_name="Hopper", email="ghopper@dlsau.edu.ph", department_id=1)
        self.room = models.Room(id=1, name="Lab 301", building="Science Bldg", capacity=40, type="lab")

        self.db.add_all([self.dept, self.sem, self.curr, self.fac, self.room])
        self.db.commit()

        self.sched = models.Schedule(
            id=1,
            semester_id=1,
            curriculum_id=1,
            faculty_id=1,
            room_id=1,
            day_of_week="Mon",
            start_time=time(8, 0),
            end_time=time(9, 30),
            section="CS-1A",
            status="draft",
            is_locked=False,
        )
        self.db.add(self.sched)
        self.db.commit()

    def tearDown(self):
        self.db.close()

    def test_attachment_header_sanitization(self):
        # DLSAU department with commas and ampersands
        res = _attachment_header("schedule_College_of_Arts,_Sciences_&_Technology", "pdf")
        self.assertIn("Content-Disposition", res)
        # Check that quotes are used and commas/special characters are replaced by safe ASCII
        val = res["Content-Disposition"]
        self.assertTrue(val.startswith('attachment; filename="'))
        self.assertTrue(val.endswith('.pdf"'))
        self.assertNotIn(",", val)
        self.assertNotIn("&", val)
        self.assertEqual(attachment_header, _attachment_header)

    def test_build_schedule_pdf_with_data_and_xml_escaping(self):
        buffer, count = build_schedule_pdf(
            [self.sched],
            dept_name="College of Arts, Sciences & Technology",
            term_label="2026-2027 1st",
            faculty_name=None,
        )
        self.assertEqual(count, 1)
        pdf_bytes = buffer.getvalue()
        self.assertTrue(pdf_bytes.startswith(b"%PDF-"), "Generated file must start with PDF magic bytes")

    def test_build_schedule_pdf_empty(self):
        buffer, count = build_schedule_pdf(
            [],
            dept_name="Empty Department",
            term_label="2026-2027 1st",
            faculty_name="Nonexistent Professor",
        )
        self.assertEqual(count, 0)
        pdf_bytes = buffer.getvalue()
        self.assertTrue(pdf_bytes.startswith(b"%PDF-"))

    def test_build_schedule_excel_with_data(self):
        buffer = build_schedule_excel([self.sched])
        excel_bytes = buffer.getvalue()
        self.assertTrue(excel_bytes.startswith(b"PK"), "Excel files are zip archives starting with PK")

        # Load back with openpyxl to verify rows and headers
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active
        self.assertIsNotNone(ws)
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)  # Header + 1 row
        self.assertEqual(rows[0][0], "Curriculum Code")
        self.assertEqual(rows[1][0], "CS101")
        self.assertEqual(rows[1][3], "Grace Hopper")

    def test_import_schedules_from_excel_valid(self):
        # Create an Excel workbook in memory with required columns
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(['CurriculumCode', 'FacultyEmail', 'RoomName', 'Day', 'StartTime', 'EndTime', 'Section'])
        ws.append(['CS101', 'ghopper@dlsau.edu.ph', 'Lab 301', 'Tue', '09:30 AM', '11:00 AM', 'CS-1B'])
        buf = io.BytesIO()
        wb.save(buf)

        success_count, errors = import_schedules_from_excel(buf.getvalue(), self.db, semester_id=1)
        self.assertEqual(success_count, 1)
        self.assertEqual(len(errors), 0)

        imported = self.db.query(models.Schedule).filter(models.Schedule.section == "CS-1B").first()
        self.assertIsNotNone(imported)
        self.assertEqual(imported.day_of_week, "Tue")

    def test_import_schedules_from_excel_missing_col_raises_400(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(['CurriculumCode', 'FacultyEmail'])  # Missing required columns
        buf = io.BytesIO()
        wb.save(buf)

        with self.assertRaises(HTTPException) as ctx:
            import_schedules_from_excel(buf.getvalue(), self.db, semester_id=1)
        self.assertEqual(ctx.exception.status_code, 400)


if __name__ == "__main__":
    unittest.main()
